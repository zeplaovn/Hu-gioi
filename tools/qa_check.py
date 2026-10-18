#!/usr/bin/env python3
"""Kiểm tra tự động gói FINAL v1.3. Chạy từ thư mục gốc gói: python3 scripts/qa_check.py"""
import csv, collections, glob, json, os, re, subprocess, sys
res = []
def check(name, ok, detail=''): res.append((name, bool(ok), detail))
def rd(f): return list(csv.DictReader(open('data/'+f, encoding='utf-8-sig')))
def txt(f): return open(f, encoding='utf-8').read()

# --- số lượng và ID
exp = {'linh_duoc.csv':72,'dan_duoc.csv':60,'nguyen_lieu.csv':54,'trang_bi.csv':82,'affix.csv':20,'set_trang_bi.csv':9,
       'quai_boss.csv':30,'drop_table.csv':132,'death_rules.csv':9,'thien_menh_cac.csv':17,'progression_60_days.csv':15,'economy_60_days.csv':60}
for f, n in exp.items(): check(f'Số dòng {f} = {n}', len(rd(f)) == n, f'thực tế {len(rd(f))}')
for f in ['linh_duoc.csv','dan_duoc.csv','nguyen_lieu.csv','trang_bi.csv','drop_table.csv','quai_boss.csv','affix.csv','set_trang_bi.csv','death_rules.csv','thien_menh_cac.csv']:
    rows = rd(f); k = list(rows[0].keys())[0]
    d = [i for i, c in collections.Counter(r[k] for r in rows).items() if c > 1]
    check(f'ID duy nhất {f}', not d, ', '.join(d[:3]))
worlds = set()
for f in ['linh_duoc.csv','dan_duoc.csv','nguyen_lieu.csv','trang_bi.csv','quai_boss.csv','economy_60_days.csv','progression_60_days.csv','shop_curve.csv']:
    for r in rd(f):
        for k, v in r.items():
            if k.lower() == 'world': worlds.add(v)
check('Chỉ ba map Hạ/Linh/Tiên', worlds == {'Hạ Giới','Linh Giới','Tiên Giới'}, str(sorted(worlds)))
prog = rd('progression_60_days.csv')
check('Tổng ngày progression = 60', sum(int(p['day_end'])-int(p['day_start']) for p in prog) == 60)

# --- kinh tế
eco = rd('economy_60_days.csv'); ok = True
for r in eco:
    s, a, d = int(r['source']), int(r['system_sink']), int(r['death_loss_expected'])
    if int(r['sink_total']) != a+d or int(r['net']) != s-(a+d) or abs(float(r['source_sink_ratio'])-round(s/(a+d), 3)) > 0.0011: ok = False
check('Economy: sink_total/net/ratio nhất quán', ok)
w7 = [sum(int(eco[i+k]['source']) for k in range(7))/sum(int(eco[i+k]['sink_total']) for k in range(7)) for i in range(len(eco)-6)]
check('Source/sink cửa sổ 7 ngày trong 0.9–1.35', all(0.9 <= x <= 1.35 for x in w7), f'min {min(w7):.3f} max {max(w7):.3f}')
check('Source/sink cửa sổ 7 ngày trong vùng khỏe 1.05–1.25', all(1.05 <= x <= 1.25 for x in w7))

# --- Thiên Mệnh
cfg = json.load(open('data/balance_config.json', encoding='utf-8'))
hm = cfg['heaven_mandate']
check('TM mỗi kiếp = tầng × giai đoạn × TM/tầng', hm['points_per_clean_life'] == hm['tang_per_realm']*hm['stages']*hm['points_per_tang'], str(hm['points_per_clean_life']))
tm = rd('thien_menh_cac.csv'); total = 0; okc = True
cmap = {'Tướng Mạo':'TuongMao','Phúc Duyên':'PhucDuyen','Căn Cốt':'CanCot','Ngộ Tính':'NgoTinh','Vận Khí':'VanKhi'}
for r in tm:
    if r['type'] == 'Permanent Stat':
        m = re.match(r'(\d+)\s*\+\s*(\d+)×level', r['cost']); b, s = int(m[1]), int(m[2]); cap = int(r['limit'])
        key = next(v for k, v in cmap.items() if k in r['name'])
        if hm['stat_costs'][key] != {'base': b, 'step': s}: okc = False
        total += sum(b+s*l for l in range(cap))
check('Giá Quán Đỉnh trong CSV khớp config', okc)
check('Tổng giá Quán Đỉnh = config', total == hm['stat_total_cost'], f'{total}')
items = [r for r in tm if r['type'] == 'Exclusive Item']
check('Giá item TM trong khoảng config', all(hm['exclusive_item_price_range'][0] <= int(r['cost']) <= hm['exclusive_item_price_range'][1] for r in items))
others = ''.join(open('data/'+f, encoding='utf-8-sig').read() for f in ['drop_table.csv','shop_curve.csv','dan_duoc.csv','trang_bi.csv','nguyen_lieu.csv','linh_duoc.csv','quai_boss.csv'])
check('Item Thiên Mệnh không xuất hiện ở drop/shop/nguồn thường', not [r['name'] for r in items if r['name'] in others])

# --- thọ nguyên
tn = rd('tho_nguyen.csv'); t = cfg['tho_nguyen']
check('tho_nguyen.csv khớp công thức', all(int(r['allowance_years']) == t['start_years']+t['grant_years_per_target_day']*int(r['day_end']) and int(r['clean_consumed_years']) == t['use_years_per_credited_day']*int(r['day_end']) for r in tn))

# --- tài liệu
mds = glob.glob('*.md')
check('Không còn ký tự xuống dòng literal (backslash-n) trong md', all(chr(92)+'n' not in txt(f) for f in mds if f not in ('CHANGELOG_v1.3.md','QA_REPORT.md')))
bad = re.compile(r'Migration save cũ|Mapping save legacy|meta\.ascended')
check('Không còn migration save cũ / meta.ascended trong tài liệu', not [f for f in mds if f not in ('CHANGELOG_v1.3.md','QA_REPORT.md') and bad.search(txt(f))])
plan = txt('Thien_Van_Dai_Luc_Ke_Hoach_Phat_Trien_FINAL.md'); bm = txt('Thien_Van_Dai_Luc_Balance_Master_60_Ngay.md')
for s in ['Phi Thăng chỉ chuyển map', '135 TM', '2.280', 'Database làm mới từ đầu', 'Thọ nguyên']:
    check(f'Kế hoạch có "{s}"', s in plan)
check('Balance Master khớp kế hoạch (135 TM, 2.280)', '135 TM' in bm and '2.280' in bm)
check('Stack/API/frontend để OPEN (không ghi công nghệ cụ thể làm baseline)', cfg.get('technology_stack') == 'OPEN' and 'OPEN' in plan)
check('Heading Balance Master không trùng số', len({m.group(1) for m in re.finditer(r'^## (\d+)\.', bm, re.M)}) == len(re.findall(r'^## \d+\.', bm, re.M)))

# --- simulator
out = subprocess.check_output([sys.executable, 'Thien_Van_Dai_Luc_Simulator_60_Ngay.py', '--sensitivity'])
saved = json.load(open('data/simulation_result_v1_3.json', encoding='utf-8'))
check('Simulator tái lập kết quả đã lưu (seed cố định)', json.loads(out) == saved)
mid = saved['sensitivity'][1]
check('Guardrail: kiếp sạch thứ 2+ median ≥ 45 ngày', mid['later_life_clean_days_median'] >= cfg['guardrails']['later_life_clean_days_min'], str(mid['later_life_clean_days_median']))
check('Phân phối thời gian không suy biến (P25 < median < P75)', mid['p25_days'] < mid['median_days'] < mid['p75_days'])

# --- xlsx
from openpyxl import load_workbook
wb = load_workbook('Thien_Van_Dai_Luc_Balance_Master_60_Ngay.xlsx', data_only=True)
errs = [c.coordinate for ws in wb for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith('#')]
check('xlsx không có lỗi công thức', not errs, ','.join(errs[:3]))
check('xlsx Dashboard khớp config (TM/kiếp, tổng giá)', wb['Dashboard']['B3'].value == hm['points_per_clean_life'] and wb['Dashboard']['B4'].value == hm['stat_total_cost'])

w = max(len(n) for n, _, _ in res)
for n, ok_, d in res: print(f"{'PASS' if ok_ else 'FAIL'}  {n:<{w}}  {d}")
sys.exit(0 if all(o for _, o, _ in res) else 1)
