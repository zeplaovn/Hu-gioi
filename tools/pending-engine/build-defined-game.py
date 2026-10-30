#!/usr/bin/env python3
"""Dựng bộ dữ liệu thiết kế theo idea/define_game.txt, không sửa luật game."""

import argparse
import csv
import json
import math
import re
import subprocess
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parent.parent
OUTPUT = ROOT / 'docs' / 'data-game-chuan'
FILES = ('vat_pham.xlsx', 'quai_boss.xlsx', 'npc.xlsx', 'linh_can.xlsx', 'du_lieu.json')
RARITIES = ['phàm', 'linh', 'huyền', 'địa', 'thiên']
REALMS = ['luyện khí', 'trúc cơ', 'kim đan', 'nguyên anh', 'hóa thần']
ELEMENTS = ['kim', 'mộc', 'thủy', 'hỏa', 'thổ', 'phong', 'độc', 'băng', 'lôi', 'hỗn độn', 'không hệ']
SOURCES = ['cửa hàng', 'rơi từ quái', 'rơi từ boss', 'nhiệm vụ', 'bí cảnh', 'sự kiện', 'chế tạo', 'npc tặng']
RELATIONS = ['sư phụ', 'đồng môn', 'bạn', 'thương nhân', 'đối thủ', 'kẻ thù', 'người lạ']
NPC_ROLES = ['bán đồ', 'giao nhiệm vụ', 'dạy bí tịnh', 'dạy công pháp', 'luyện đan thuê']
NODE_DATA = r"""
const items = require('./shared/data/items.js');
const {Game} = require('./shared/game.js');
const {DUNGEONS} = require('./shared/data/adventures.js');
const {ZONES} = require('./shared/data/catalog.js');
const game = new Game();
const selected = items.filter(x => x.tier <= 4);
console.log(JSON.stringify({items:selected, recipes:selected.filter(x=>x.kind==='PILL').map(x=>game.recipe(x.code)),
 dungeons:DUNGEONS.filter(x=>x.realm<=4), zones:ZONES.filter(x=>x.realm<=4)}));
"""


def source_rows(path, name):
    wb = load_workbook(ROOT / path, read_only=True, data_only=True)
    try:
        return list(wb[name].values)[6:]
    finally:
        wb.close()


def mapped_id(prefix, old):
    match = re.search(r'(\d{3})$', old)
    if not match:
        raise ValueError('Không thể chuyển mã: ' + old)
    return prefix + match.group(1)


def element(name):
    text = name.lower()
    for keyword, value in [('hỗn độn', 'hỗn độn'), ('băng', 'băng'), ('hàn', 'băng'), ('tuyết', 'băng'),
                           ('lôi', 'lôi'), ('điện', 'lôi'), ('phong', 'phong'), ('độc', 'độc'),
                           ('kim', 'kim'), ('mộc', 'mộc'), ('thủy', 'thủy'), ('hỏa', 'hỏa'),
                           ('viêm', 'hỏa'), ('thổ', 'thổ')]:
        if keyword in text:
            return value
    return 'không hệ'


def add_source(row, source):
    values = row['nguồn gốc'].split('; ')
    if source not in values:
        values.append(source)
    row['nguồn gốc'] = '; '.join(values)


def build_data(raw, enemy_assets, npc_assets, portraits):
    items = raw['items']
    data = {key: [] for key in ['hạt giống', 'linh thảo', 'đan dược', 'đan phương', 'pháp bảo',
                                 'hiệu ứng pháp bảo', 'công pháp', 'bí tịch', 'rèn pháp bảo', 'phù lục', 'vẽ phù',
                                 'quái vật', 'boss', 'giai đoạn boss', 'vật phẩm rơi', 'npc', 'linh căn',
                                 'loại linh căn', 'ngũ hành', 'quy tắc linh căn', 'quy tắc npc']}
    by_old = {x['code']: x for x in items}
    herbs = {x['code']: x for x in items if x['kind'] == 'HERB'}
    seed_by_herb = {x['herb']: mapped_id('HG', x['code']) for x in items if x['kind'] == 'SEED'}
    times, yields, soils = [10, 30, 120, 360, 1440], [(3, 5), (2, 4), (1, 3), (1, 2), (1, 1)], [
        'đất thường', 'linh điền', 'linh điền', 'linh điền', 'linh điền có trận pháp']

    for x in sorted((x for x in items if x['kind'] == 'SEED'), key=lambda x: x['code']):
        herb = herbs[x['herb']]
        tier = herb['tier']
        lo, hi = yields[tier]
        soil = 'hàn đàm' if element(herb['name']) == 'băng' else 'hỏa địa' if element(herb['name']) == 'hỏa' else soils[tier]
        data['hạt giống'].append({'mã hạt': mapped_id('HG', x['code']), 'tên hạt': x['name'], 'nguồn gốc': 'cửa hàng',
            'độ khó': tier + 1, 'thời gian trồng (phút)': times[tier], 'linh thảo thu được': mapped_id('LT', x['herb']),
            'sản lượng tối thiểu': lo, 'sản lượng tối đa': hi, 'điều kiện đất': soil, 'mã game cũ': x['code']})
    recipes = {r['code']: r for r in raw['recipes']}
    for x in sorted(herbs.values(), key=lambda x: x['code']):
        used = any(x['code'] in r['herbs'] for r in recipes.values())
        direct = x['code'] == 'HERB-003'
        data['linh thảo'].append({'mã cây': mapped_id('LT', x['code']), 'tên cây': x['name'], 'nguồn gốc': 'cửa hàng',
            'trồng từ hạt': seed_by_herb.get(x['code']), 'độ hiếm': RARITIES[x['tier']], 'tuổi dược (năm)': x['years'],
            'hệ': element(x['name']), 'cách dùng': ('ăn trực tiếp; ' if direct else '') + ('luyện đan; bán' if used else 'bán'),
            'hiệu quả': ('hồi 50 hp khi ăn; ' if direct else '') + ('nguyên liệu đan phương; bán' if used else 'bán'), 'mã game cũ': x['code']})
    for x in sorted((x for x in items if x['kind'] == 'PILL'), key=lambda x: x['code']):
        code = mapped_id('DD', x['code'])
        data['đan dược'].append({'mã đan': code, 'tên đan': x['name'], 'độ hiếm': RARITIES[x['tier']],
            'yêu cầu': REALMS[x['tier']], 'nguồn gốc': 'chế tạo', 'mã hiệu ứng game cũ': x['effect'],
            'mô tả tham khảo': x['description'], 'mã game cũ': x['code']})
        for herb, count in recipes[x['code']]['herbs'].items():
            if herb not in herbs:
                raise ValueError('Nguyên liệu nằm ngoài phạm vi 5 cảnh giới: ' + herb)
            data['đan phương'].append({'mã đan': code, 'mã linh thảo': mapped_id('LT', herb), 'số lượng': count})

    effects = ['giáp ảo', 'hồi máu', 'đóng băng', 'gây độc', 'chảy máu', 'gây choáng']
    for x in sorted((x for x in items if x['kind'] == 'EQ'), key=lambda x: x['code']):
        t, code = x['tier'], mapped_id('PB', x['code'])
        name = x['name'].lower()
        kind = 'vũ khí' if any(w in name for w in ('kiếm', 'đao')) else 'giáp' if any(w in name for w in ('pháp y', 'linh ngoa')) else 'trang sức' if any(w in name for w in ('linh giới', 'tâm bội')) else 'pháp khí hỗ trợ'
        data['pháp bảo'].append({'mã PB': code, 'tên pháp bảo': x['name'], 'độ hiếm': RARITIES[t], 'loại': kind,
            'công dụng': effects[(int(code[-3:]) - 1) % 6], 'hp cộng': [20, 70, 160, 350, 700][t],
            'mp cộng': [5, 15, 35, 70, 140][t], 'atk cộng': [8, 18, 35, 65, 120][t],
            'def cộng': [3, 8, 18, 40, 80][t], 'eva cộng (%)': t + 1, 'sta cộng': [5, 10, 20, 40, 80][t],
            'crit cộng (%)': t + 1, 'crit_dame cơ bản (%)': 150, 'yêu cầu': REALMS[t],
            'nguồn gốc': 'cửa hàng; chế tạo', 'mã game cũ': x['code']})
        effect = effects[(int(code[-3:]) - 1) % 6]
        params = {'mã PB': code, 'hiệu ứng': effect, 'điều kiện kích hoạt': 'khi kích hoạt' if effect in ('giáp ảo', 'hồi máu') else 'khi đánh trúng',
                  'hồi chiêu (giây)': 15 + 2*t if effect in ('giáp ảo', 'hồi máu') else 0,
                  'tỉ lệ kích hoạt (%)': 100, 'thời gian (giây)': 3 + t,
                  'lá chắn (% hp tối đa)': None, 'hồi máu (hp)': None, 'chu kỳ hồi (giây)': None,
                  'độc (sát thương/giây)': None, 'số lớp tối đa': None, 'chảy máu (% hp hiện tại/giây)': None}
        if effect == 'giáp ảo': params['lá chắn (% hp tối đa)'] = 10 + 2*t
        elif effect == 'hồi máu': params.update({'hồi máu (hp)': 10*(t+1), 'chu kỳ hồi (giây)': 1})
        elif effect in ('đóng băng', 'gây choáng'): params.update({'tỉ lệ kích hoạt (%)': 10 + t, 'thời gian (giây)': 1 + t * .25})
        elif effect == 'gây độc': params.update({'độc (sát thương/giây)': 5*(t+1), 'số lớp tối đa': 3})
        else: params['chảy máu (% hp hiện tại/giây)'] = 1 + t
        data['hiệu ứng pháp bảo'].append(params)

    # Chỉ dùng 20 công pháp đầu có loại phù hợp tên; các BT cao hơn của game
    # có tên và effect không nhất quán nên không tự suy diễn thành kỹ năng.
    base_manuals = [by_old[f'BT-{i:03d}'] for i in range(1, 21)]
    for x in base_manuals:
        t, old = x['tier'], x['code']
        if x['effect'] in ('SAT', 'PHONG'):
            attack = x['effect'] == 'SAT'
            data['công pháp'].append({'mã CP': mapped_id('CP', old), 'tên công pháp': x['name'], 'độ hiếm': RARITIES[t],
                'hệ': element(x['name']), 'mô tả': x['description'], 'công dụng': 'công kích' if attack else 'phòng ngự',
                'tiêu hao mp': 10 + 5*t, 'tiêu hao sta': 0, 'hồi chiêu (giây)': 5 + 2*t,
                'atk_multiplier': round(1.2 + .1*t, 2) if attack else None, 'def_bonus': None if attack else 15*(t+1),
                'heal': None, 'buff stat': None, 'buff giá trị': None, 'thời gian buff (giây)': None if attack else 6+t,
                'yêu cầu': REALMS[t], 'nguồn gốc': 'cửa hàng', 'mã game cũ': old})
        else:
            cultivation = x['effect'] == 'TOC'
            data['bí tịch'].append({'mã BT': mapped_id('BT', old), 'tên bí tịch': x['name'],
                'thể loại': 'luyện khí' if cultivation else 'luyện đan', 'độ hiếm': RARITIES[t],
                'công dụng': f'tăng {10+5*t}% tốc độ tu luyện' if cultivation else f'tăng {5+2*t} điểm % tỉ lệ thành đan',
                'yêu cầu học': REALMS[t], 'nguồn gốc': 'cửa hàng', 'mã game cũ': old})
    for code, name, role, rare, desc in [
        ('BT900', 'Cơ Bản Luyện Bảo', 'luyện pháp bảo', 'linh', 'mở các công thức rèn pháp bảo trong bảng rèn'),
        ('BT901', 'Hỏa Phù Chân Giải', 'vẽ phù lục', 'linh', 'mở công thức phù PL001')]:
        data['bí tịch'].append({'mã BT': code, 'tên bí tịch': name, 'thể loại': role, 'độ hiếm': rare,
            'công dụng': desc, 'yêu cầu học': 'luyện khí', 'nguồn gốc': 'cửa hàng', 'mã game cũ': None})
    for code, name, role, el, cost, cd, heal, buff_stat, buff_value, duration in [
        ('CP900', 'Xuân Hồi Quyết', 'support', 'mộc', 25, 10, 120, None, None, None),
        ('CP901', 'Cuồng Phong Bộ', 'support', 'phong', 30, 20, None, 'eva', 15, 10)]:
        data['công pháp'].append({'mã CP': code, 'tên công pháp': name, 'độ hiếm': 'linh', 'hệ': el,
            'mô tả': 'hồi máu' if heal else 'tăng né tránh', 'công dụng': role, 'tiêu hao mp': cost, 'tiêu hao sta': 0,
            'hồi chiêu (giây)': cd, 'atk_multiplier': None, 'def_bonus': None, 'heal': heal, 'buff stat': buff_stat,
            'buff giá trị': buff_value, 'thời gian buff (giây)': duration, 'yêu cầu': 'trúc cơ',
            'nguồn gốc': 'cửa hàng', 'mã game cũ': None})
    data['phù lục'].append({'mã phù': 'PL001', 'tên phù': 'Hỏa Phù', 'hệ': 'hỏa', 'hiệu quả': 'gây 120% ATK lên mục tiêu',
                             'bí tịch yêu cầu': 'BT901', 'nguồn gốc': 'chế tạo', 'mã game cũ': None})
    data['vẽ phù'].append({'mã phù': 'PL001', 'mã bí tịch': 'BT901', 'mã linh thảo': 'LT001', 'số lượng': 1})
    data['linh thảo'][0]['cách dùng'] += '; vẽ phù lục'
    data['linh thảo'][0]['hiệu quả'] += '; nguyên liệu vẽ PL001'
    for i, pb in enumerate(data['pháp bảo']):
        herb = data['linh thảo'][i % len(data['linh thảo'])]
        data['rèn pháp bảo'].append({'mã PB': pb['mã PB'], 'mã bí tịch': 'BT900', 'mã linh thảo': herb['mã cây'], 'số lượng': 2})
        if 'luyện pháp bảo' not in herb['cách dùng']:
            herb['cách dùng'] += '; luyện pháp bảo'
            herb['hiệu quả'] += '; nguyên liệu rèn pháp bảo'

    mobs = [r for r in enemy_assets if r and r[0] and r[2] == 'Hạ Giới' and r[3] == 'Enemy']
    for i, r in enumerate(mobs):
        t = min(i, 4)
        data['quái vật'].append({'mã quái': f'QV{i+1:03}', 'tên quái': r[1], 'cấp': [1, 5, 10, 15, 20][t],
            'hệ': ['không hệ', 'thổ', 'phong', 'thủy', 'kim'][t], 'khu vực': raw['zones'][t]['name'],
            'hp': [100, 300, 650, 1300, 2500][t], 'mp': [0, 20, 40, 80, 150][t], 'atk': [8, 20, 40, 75, 120][t],
            'def': [2, 8, 18, 35, 60][t], 'eva (%)': 3+t, 'sta': 20+20*t, 'crit (%)': 2+t,
            'crit_dame (%)': 150, 'kỹ năng': ['cắn: chảy máu 1% hp hiện tại/giây trong 3 giây', 'vồ: 120% ATK',
                'gió lốc: 130% ATK', 'quật đuôi: 140% ATK', 'kim vũ: 150% ATK'][t], 'exp': [15, 60, 150, 350, 700][t],
            'mã ảnh cũ': r[0]})
    boss_info = [('Cáo Truyền Thuyết', 4, 'không hệ', 'Hóa Thần Đài', 'BOSS-HG-FINAL-001')]
    dungeon_elements = ['mộc', 'thủy', 'hỏa', 'kim', 'thổ', 'không hệ']
    for i, d in enumerate(raw['dungeons']):
        boss_info.append((d['boss'], d['realm'], dungeon_elements[i], d['name'], None))
    for i, (name, tier, el, area, legacy) in enumerate(boss_info):
        code = f'BS{i+1:03}'
        data['boss'].append({'mã boss': code, 'tên boss': name, 'cấp': 10 + tier*10, 'hệ': el, 'khu vực': area,
            'hp': 1500 * (tier+1)**2, 'mp': 150*(tier+1), 'atk': 35*(tier+1), 'def': 20*(tier+1),
            'eva (%)': 3+tier, 'sta': 100*(tier+1), 'crit (%)': 5+tier, 'crit_dame (%)': 180,
            'kỹ năng': 'trọng kích; linh lực bộc phát; cuồng nộ', 'exp': 300*(tier+1)**2, 'mã ảnh cũ': legacy})
        for start, end, skills in [(100, 50, 'trọng kích 130% ATK'), (50, 20, 'linh lực bộc phát 160% ATK'),
                                    (20, 0, 'cuồng nộ 200% ATK')]:
            data['giai đoạn boss'].append({'mã boss': code, 'hp từ (%)': start, 'hp đến (%)': end, 'kỹ năng giai đoạn': skills})
    loot_items = {row[key]: row for sheet, key in [('hạt giống','mã hạt'), ('linh thảo','mã cây'), ('pháp bảo','mã PB'),
                                                     ('công pháp','mã CP'), ('bí tịch','mã BT')] for row in data[sheet]}
    for i, mob in enumerate(data['quái vật']):
        for sheet, key, rate in [('hạt giống', 'mã hạt', 30), ('linh thảo', 'mã cây', 20), ('pháp bảo', 'mã PB', 2)]:
            item = data[sheet][i * len(data[sheet]) // len(data['quái vật'])][key]
            data['vật phẩm rơi'].append({'mã đối thủ': mob['mã quái'], 'mã vật phẩm': item, 'tỉ lệ (%)': rate, 'số lượng tối thiểu': 1, 'số lượng tối đa': 1})
            add_source(loot_items[item], 'rơi từ quái')
    for i, boss in enumerate(data['boss']):
        for sheet, key, rate in [('pháp bảo', 'mã PB', 15), ('công pháp', 'mã CP', 5), ('bí tịch', 'mã BT', 3)]:
            item = data[sheet][i % len(data[sheet])][key]
            data['vật phẩm rơi'].append({'mã đối thủ': boss['mã boss'], 'mã vật phẩm': item, 'tỉ lệ (%)': rate, 'số lượng tối thiểu': 1, 'số lượng tối đa': 1})
            add_source(loot_items[item], 'rơi từ boss')

    profiles = {
        'NPC-01': ('Người quen đường quanh Trấn Lạc Phong.', 'người lạ', 'giao nhiệm vụ'),
        'NPC-02': ('Kiếm tu của Thanh Vân Tông.', 'đồng môn', 'dạy công pháp'),
        'NPC-03': ('Tu sĩ Huyết Sát Môn, thích tỉ thí.', 'đối thủ', 'giao nhiệm vụ'),
        'NPC-04': ('Dược sư của Vạn Dược Cốc.', 'đồng môn', 'luyện đan thuê'),
        'NPC-05': ('Người tìm hiểu di tích Huyền Quy.', 'người lạ', 'giao nhiệm vụ'),
        'NPC-10': ('Người quen đường trong Hắc Lâm.', 'bạn', 'giao nhiệm vụ'),
        'NPC-11': ('Người thợ ở Mỏ Linh Thiết.', 'bạn', 'bán đồ'),
        'NPC-12': ('Thương nhân Lạc Phong Phường Thị.', 'thương nhân', 'bán đồ'),
        'NPC-13': ('Tán tu khảo sát Suối Địa Mạch.', 'người lạ', 'dạy bí tịnh'),
        'NPC-14': ('Lão nhân ẩn cư tại Kiếm Cốt Lâm.', 'sư phụ', 'dạy công pháp'),
    }
    for i, r in enumerate(row for row in npc_assets if row and row[0] in profiles):
        old, name, world, area = r[:4]
        background, relation, role = profiles[old]
        chosen = (data['công pháp'][i % len(data['công pháp'])]['mã CP'] if role == 'dạy công pháp' else
                  data['bí tịch'][i % len(data['bí tịch'])]['mã BT'] if role == 'dạy bí tịnh' else
                  data['hạt giống'][i % len(data['hạt giống'])]['mã hạt'] if role == 'bán đồ' else
                  data['đan dược'][i % len(data['đan dược'])]['mã đan'] if role == 'luyện đan thuê' else None)
        path = 'media/assets/chars/npc/' + portraits[old]
        data['npc'].append({'mã NPC': f'NPC{int(old[-2:]):03}', 'tên NPC': name, 'bối cảnh': background,
            'quan hệ': relation, 'thích (mã vật phẩm)': data['linh thảo'][i]['mã cây'],
            'không thích (mã vật phẩm)': data['pháp bảo'][i]['mã PB'],
            'không thích khác': ['gian dối', 'phá hoại', 'thất hứa'][i % 3], 'thiện cảm ban đầu': 0,
            'vai trò': role, 'mã hàng/dạy/đan': chosen, 'địa điểm': area, 'giới': world, 'ảnh': path, 'mã ảnh cũ': old})

    root_specs = [
        ('LC001','Kim Linh Căn','ngũ hành','kim','phàm','bẩm sinh',1,'atk +5%; crit +2%','thủy','mộc','công pháp kim +10% sát thương', ''),
        ('LC002','Mộc Linh Căn','ngũ hành','mộc','phàm','bẩm sinh',1,'hp +8%; hồi máu +5%','hỏa','thổ','hiệu quả hồi máu +10%', ''),
        ('LC003','Thủy Linh Căn','ngũ hành','thủy','phàm','bẩm sinh',1,'mp +8%; eva +2%','mộc','hỏa','hồi mp +10%', ''),
        ('LC004','Hỏa Linh Căn','ngũ hành','hỏa','phàm','bẩm sinh',1,'atk +8%; crit_dame +5%','thổ','kim','công pháp hỏa +10% sát thương', ''),
        ('LC005','Thổ Linh Căn','ngũ hành','thổ','phàm','bẩm sinh',1,'def +8%; hp +5%','kim','thủy','giảm 5% sát thương nhận', ''),
        ('LC101','Phong Linh Căn','dị','phong','huyền','biến dị từ mộc hoặc thủy',2,'eva +10%; sta +10%','', 'thổ','tốc di chuyển và ra chiêu +10%','mộc; thủy'),
        ('LC102','Độc Linh Căn','dị','độc','huyền','biến dị từ mộc hoặc thổ',2,'atk +5%; hp +5%','', 'mộc','15% gây độc; kháng độc 100%','mộc; thổ'),
        ('LC103','Băng Linh Căn','dị','băng','huyền','biến dị từ thủy',2,'def +8%; mp +10%','', 'hỏa','10% đóng băng 1,5 giây','thủy'),
        ('LC104','Lôi Linh Căn','dị','lôi','huyền','biến dị từ hỏa hoặc kim',2,'atk +10%; crit +5%','', 'thủy','10% choáng 1 giây; xuyên 10% def','hỏa; kim'),
        ('LC201','Hỗn Độn Linh Căn','hỗn độn','hỗn độn','thiên','biến dị sau thiên kiếp',3,
         'hp +15%; mp +15%; atk +10%; def +10%; eva +5%; crit +5%; crit_dame +20%', '', 'ngũ hành và dị',
         'không bị khắc; học mọi công pháp; hệ thành phần +20%', ''),
    ]
    for code, name, group, el, rarity, origin, rate, stats, generates, counters, special, parents in root_specs:
        data['linh căn'].append({'mã LC': code, 'tên linh căn': name, 'loại': group, 'hệ': el, 'độ hiếm': rarity,
            'nguồn gốc': origin, 'tốc độ tu luyện cơ sở': rate, 'stat cộng thêm': stats, 'tương sinh': generates,
            'khắc theo ví dụ': counters, 'hệ gốc dị': parents, 'hiệu ứng đặc trưng': special})
    for n, label, rate in [(1,'đơn linh căn',1.5),(2,'song linh căn',1.2),(3,'tam linh căn',1),(4,'tứ linh căn',.8),(5,'ngũ linh căn',.6)]:
        data['loại linh căn'].append({'loại': label, 'số hệ': n, 'hệ số tốc tu': rate, 'độ tinh khiết tối thiểu (%)': 0, 'độ tinh khiết tối đa (%)': 100})
    for el, gen, counter in [('mộc','hỏa','thổ'),('hỏa','thổ','kim'),('thổ','kim','thủy'),('kim','thủy','mộc'),('thủy','mộc','hỏa')]:
        data['ngũ hành'].append({'hệ': el, 'sinh': gen, 'khắc': counter, 'thưởng sát thương khi khắc (%)': 20})
    data['quy tắc linh căn'] = [
        {'quy tắc': 'Độ tinh khiết', 'giá trị': '0–100%; ảnh hưởng hiệu quả công pháp cùng hệ; công thức cụ thể chưa được đặc tả.'},
        {'quy tắc': 'Tốc tu ngũ hành', 'giá trị': 'Hệ số theo số hệ tại bảng Loại linh căn; hệ số cơ sở LC001–LC005 là 1.'},
        {'quy tắc': 'Tốc tu dị/hỗn độn', 'giá trị': 'Dị ×2; Hỗn Độn ×3. Không tự nhân thêm hệ số số hệ khi chưa được xác nhận.'},
        {'quy tắc': 'Dị linh căn khắc hệ', 'giá trị': 'Quy tắc chung nói khắc hệ gốc; ví dụ ghi ở cột khắc theo ví dụ. Hai cách mâu thuẫn, cần quyết định.'},
        {'quy tắc': 'Hỗn Độn: hình thành', 'giá trị': 'Từ ngũ hành có 1–5 hệ sau thiên kiếp lên đại cảnh giới; xác suất cơ bản tham khảo 1%, không bẩm sinh hay đoạt lấy.'},
        {'quy tắc': 'Hỗn Độn: thành phần', 'giá trị': 'Giữ các hệ ngũ hành trước biến dị; không bị khắc, khắc ngũ hành và dị +10% sát thương.'},
        {'quy tắc': 'Hỗn Độn: công pháp', 'giá trị': 'Học mọi hệ; hệ thành phần +20% hiệu quả; mỗi lần đột phá đại cảnh giới thêm 1 hiệu ứng đặc trưng.'},
        {'quy tắc': 'Công pháp khác hệ', 'giá trị': 'Cùng hệ được cộng, khác hệ bị trừ; mức cụ thể chưa được đặc tả. Hỗn Độn và không hệ là ngoại lệ.'},
        {'quy tắc': 'Tuổi dược', 'giá trị': 'Tuổi dược ảnh hưởng chất lượng đan; công thức cụ thể chưa được đặc tả.'},
    ]
    data['quy tắc npc'] = [
        {'quy tắc': 'Thiện cảm', 'giá trị': 'Từ 0 đến 100; khi tặng đúng món thích +5, món không thích −5; giới hạn trong khoảng 0–100. Mức ±5 là đề xuất.'},
        {'quy tắc': 'Quan hệ và vai trò', 'giá trị': 'Quan hệ và vai trò là dữ liệu thiết kế; chưa được nối vào gameplay hiện tại.'},
    ]
    data['metadata'] = {'trạng thái': 'thiết kế đề xuất; chưa tích hợp gameplay', 'phạm vi': '5 cảnh giới đầu theo đặc tả',
        'cảnh giới': REALMS, 'độ hiếm': RARITIES, 'hệ': ELEMENTS, 'nguồn gốc vật phẩm': SOURCES,
        'mã mở rộng': {'DD': 'đan dược', 'PL': 'phù lục'},
        'lưu ý dị linh căn': 'Mâu thuẫn: quy tắc chung nói khắc hệ gốc, nhưng ví dụ LC101/103/104 khắc hệ khác; lưu cả hai và chờ quyết định.',
        'lưu ý hỗn độn': 'Đặc tả cho phép mọi số hệ ngũ hành; game hiện tại chỉ cho Ngụy sau lôi kiếp. Chưa thay luật game.',
        'thiện cảm tặng thích': 5, 'thiện cảm tặng ghét': -5}
    validate(data)
    return data


def validate(data):
    def require(ok, message):
        if not ok:
            raise ValueError(message)

    primary = [('hạt giống','mã hạt','HG'), ('linh thảo','mã cây','LT'), ('đan dược','mã đan','DD'),
               ('pháp bảo','mã PB','PB'), ('công pháp','mã CP','CP'), ('bí tịch','mã BT','BT'),
               ('phù lục','mã phù','PL'), ('quái vật','mã quái','QV'), ('boss','mã boss','BS'),
               ('npc','mã NPC','NPC'), ('linh căn','mã LC','LC')]
    ids = {}
    for sheet, key, prefix in primary:
        values = [r[key] for r in data[sheet]]
        if len(values) != len(set(values)) or any(not re.fullmatch(prefix + r'\d{3}', v) for v in values):
            raise ValueError('Mã trùng hoặc sai định dạng: ' + sheet)
        ids[sheet] = set(values)
    all_items = set().union(*(ids[s] for s in ('hạt giống','linh thảo','đan dược','pháp bảo','công pháp','bí tịch','phù lục')))
    for sheet in ('hạt giống','linh thảo','đan dược','pháp bảo','công pháp','bí tịch','phù lục'):
        for r in data[sheet]:
            if not set(r['nguồn gốc'].split('; ')) <= set(SOURCES): raise ValueError('Nguồn gốc sai: ' + str(r))
    for r in data['hạt giống']:
        require(r['linh thảo thu được'] in ids['linh thảo'] and 1 <= r['độ khó'] <= 5, 'Hạt giống không hợp lệ: ' + str(r))
        require(r['thời gian trồng (phút)'] > 0 and 0 < r['sản lượng tối thiểu'] <= r['sản lượng tối đa'], 'Thời gian/sản lượng sai: ' + str(r))
    for r in data['linh thảo']:
        require(not r['trồng từ hạt'] or r['trồng từ hạt'] in ids['hạt giống'], 'Hạt của linh thảo không tồn tại: ' + str(r))
        require(r['tuổi dược (năm)'] > 0, 'Tuổi dược không hợp lệ: ' + str(r))
    herb_seeds = {r['mã cây']: r['trồng từ hạt'] for r in data['linh thảo']}
    for r in data['hạt giống']:
        require(herb_seeds[r['linh thảo thu được']] == r['mã hạt'], 'Liên kết hạt–cây không đối xứng: ' + str(r))
    for r in data['đan phương']:
        require(r['mã đan'] in ids['đan dược'] and r['mã linh thảo'] in ids['linh thảo'] and r['số lượng'] > 0, 'Đan phương không hợp lệ: ' + str(r))
    require({r['mã đan'] for r in data['đan phương']} == ids['đan dược'], 'Thiếu công thức đan')
    for r in data['rèn pháp bảo']:
        require(r['mã PB'] in ids['pháp bảo'] and r['mã bí tịch'] in ids['bí tịch'] and r['mã linh thảo'] in ids['linh thảo'] and r['số lượng'] > 0, 'Công thức rèn không hợp lệ: ' + str(r))
    for r in data['vẽ phù']:
        require(r['mã phù'] in ids['phù lục'] and r['mã bí tịch'] in ids['bí tịch'] and r['mã linh thảo'] in ids['linh thảo'] and r['số lượng'] > 0, 'Công thức phù không hợp lệ: ' + str(r))
    for r in data['phù lục']:
        require(r['bí tịch yêu cầu'] in ids['bí tịch'] and r['hệ'] in ELEMENTS, 'Phù lục không hợp lệ: ' + str(r))
    for r in data['hiệu ứng pháp bảo']:
        require(r['mã PB'] in ids['pháp bảo'] and 0 <= r['tỉ lệ kích hoạt (%)'] <= 100 and r['thời gian (giây)'] > 0 and r['hồi chiêu (giây)'] >= 0, 'Hiệu ứng không hợp lệ: ' + str(r))
    for code in ids['pháp bảo']:
        require(1 <= sum(r['mã PB'] == code for r in data['hiệu ứng pháp bảo']) <= 2, 'Pháp bảo cần 1–2 hiệu ứng: ' + code)
    for r in data['vật phẩm rơi']:
        require(r['mã đối thủ'] in ids['quái vật'] | ids['boss'] and r['mã vật phẩm'] in all_items and 0 <= r['tỉ lệ (%)'] <= 100 and 0 < r['số lượng tối thiểu'] <= r['số lượng tối đa'], 'Vật phẩm rơi không hợp lệ: ' + str(r))
    drops = {(r['mã vật phẩm'], 'rơi từ quái' if r['mã đối thủ'] in ids['quái vật'] else 'rơi từ boss') for r in data['vật phẩm rơi']}
    item_rows = {r[key]: r for sheet, key, _ in primary if sheet in ('hạt giống','linh thảo','đan dược','pháp bảo','công pháp','bí tịch','phù lục') for r in data[sheet]}
    for code, source in drops:
        require(source in item_rows[code]['nguồn gốc'].split('; '), 'Thiếu nguồn rơi: ' + code)
    for code, r in item_rows.items():
        for source in ('rơi từ quái', 'rơi từ boss'):
            require(source not in r['nguồn gốc'].split('; ') or (code, source) in drops, 'Nguồn rơi không có bảng rơi: ' + code)
    for r in data['giai đoạn boss']:
        require(r['mã boss'] in ids['boss'] and 0 <= r['hp đến (%)'] < r['hp từ (%)'] <= 100, 'Giai đoạn boss không hợp lệ: ' + str(r))
    for code in ids['boss']:
        phases = sorted((r for r in data['giai đoạn boss'] if r['mã boss'] == code), key=lambda r: -r['hp từ (%)'])
        require(bool(phases) and phases[0]['hp từ (%)'] == 100 and phases[-1]['hp đến (%)'] == 0
                and all(a['hp đến (%)'] == b['hp từ (%)'] for a, b in zip(phases, phases[1:])), 'Giai đoạn boss không liên tục: ' + code)
    for r in data['npc']:
        require(r['quan hệ'] in RELATIONS and r['vai trò'] in NPC_ROLES and 0 <= r['thiện cảm ban đầu'] <= 100, 'Quan hệ NPC không hợp lệ: ' + str(r))
        require(r['thích (mã vật phẩm)'] in all_items and r['không thích (mã vật phẩm)'] in all_items
                and r['thích (mã vật phẩm)'] != r['không thích (mã vật phẩm)'], 'Sở thích NPC không hợp lệ: ' + str(r))
        require(not r['mã hàng/dạy/đan'] or r['mã hàng/dạy/đan'] in all_items, 'Vật phẩm NPC không tồn tại: ' + str(r))
        expected = {'dạy công pháp': 'công pháp', 'dạy bí tịnh': 'bí tịch', 'luyện đan thuê': 'đan dược'}
        require(r['vai trò'] not in expected or r['mã hàng/dạy/đan'] in ids[expected[r['vai trò']]], 'Nội dung NPC không khớp vai trò: ' + str(r))
        require((ROOT / r['ảnh']).is_file(), 'Ảnh NPC không tồn tại: ' + str(r))
    for sheet in ('linh thảo','đan dược','pháp bảo','công pháp','bí tịch','linh căn'):
        require(all(r['độ hiếm'] in RARITIES for r in data[sheet]), 'Độ hiếm không hợp lệ: ' + sheet)
    for sheet in ('linh thảo','công pháp','quái vật','boss','linh căn'):
        require(all(r['hệ'] in ELEMENTS for r in data[sheet]), 'Hệ không hợp lệ: ' + sheet)
    for sheet in ('đan dược','pháp bảo','công pháp'):
        require(all(r['yêu cầu'] in REALMS for r in data[sheet]), 'Cảnh giới không hợp lệ: ' + sheet)
    require(all(r['yêu cầu học'] in REALMS for r in data['bí tịch']), 'Cảnh giới học bí tịch không hợp lệ')
    for r in data['công pháp']:
        require(r['công dụng'] in ('công kích', 'phòng ngự', 'support') and r['tiêu hao mp'] >= 0
                and r['tiêu hao sta'] >= 0 and r['hồi chiêu (giây)'] >= 0, 'Công pháp không hợp lệ: ' + str(r))
        require((r['công dụng'] != 'công kích' or (r['atk_multiplier'] is not None and r['atk_multiplier'] > 0))
                and (r['công dụng'] != 'phòng ngự' or (r['def_bonus'] is not None and r['thời gian buff (giây)'] > 0))
                and (r['công dụng'] != 'support' or (r['heal'] is not None or r['buff stat'] is not None)), 'Thiếu chỉ số công pháp: ' + str(r))
    for sheet in ('quái vật', 'boss'):
        for r in data[sheet]:
            require(r['cấp'] > 0 and r['hp'] > 0 and r['mp'] >= 0 and r['atk'] > 0 and r['def'] >= 0 and r['sta'] >= 0
                    and 0 <= r['eva (%)'] <= 100 and 0 <= r['crit (%)'] <= 100 and r['crit_dame (%)'] >= 100
                    and r['exp'] > 0 and bool(r['kỹ năng']), 'Chỉ số đối thủ không hợp lệ: ' + str(r))


def add_table(wb, name, rows):
    ws = wb.create_sheet(name)
    headers = list(rows[0]) if rows else []
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h) for h in headers])
    ws.freeze_panes = 'C2'
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    ws.row_dimensions[1].height = 30
    for cell in ws[1]:
        cell.fill = PatternFill('solid', fgColor='27485B')
        cell.font = Font(bold=True, color='FFFFFF')
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
            if cell.row % 2 == 0: cell.fill = PatternFill('solid', fgColor='F2F7F8')
    for col, header in enumerate(headers, 1):
        width = min(48, max(12, len(header) + 2, *(len(str(r.get(header) or '')) + 2 for r in rows)))
        ws.column_dimensions[get_column_letter(col)].width = width
    for i, row in enumerate(rows, 2):
        lines = max((math.ceil(len(str(v)) / (ws.column_dimensions[get_column_letter(j)].width - 2))
                     for j, v in enumerate(row.values(), 1) if v is not None), default=1)
        ws.row_dimensions[i].height = min(90, max(18, lines * 15))


def make_book(data, sheets):
    wb = Workbook()
    wb.remove(wb.active)
    notes = [
        {'chủ đề':'Trạng thái', 'giải thích':'Bản thiết kế đề xuất theo define_game.txt, chưa thay dữ liệu hoặc luật đang chạy.'},
        {'chủ đề':'Phạm vi', 'giải thích':'Chỉ 5 cảnh giới đầu của đặc tả; tên/mã game cũ để đối chiếu. Kim Đan tương ứng Kết Đan trong game.'},
        {'chủ đề':'Đơn vị', 'giải thích':'Thời gian trồng tính bằng phút thực; hồi chiêu và hiệu ứng tính bằng giây; tỉ lệ phần trăm ghi 0–100.'},
        {'chủ đề':'Pháp bảo', 'giải thích':'crit_dame 150 là giá trị cơ bản, không cộng 150% cho từng món. Các chỉ số và hiệu ứng là đề xuất.'},
        {'chủ đề':'Vật phẩm rơi', 'giải thích':'Mỗi dòng là một lần xét độc lập; tỉ lệ 0–100%. Sản lượng là số nguyên trong khoảng tối thiểu–tối đa.'},
        {'chủ đề':'Mã mở rộng', 'giải thích':'DD = đan dược, PL = phù lục; bổ sung để liên kết đan phương và bí tịch trong đặc tả.'},
        {'chủ đề':'Mâu thuẫn linh căn', 'giải thích': data['metadata']['lưu ý dị linh căn']},
    ]
    add_table(wb, 'Hướng dẫn', notes)
    wb['Hướng dẫn'].column_dimensions['B'].width = 100
    for sheet, title in sheets:
        add_table(wb, title, data[sheet])
    return wb


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', type=Path, default=OUTPUT)
    parser.add_argument('--force', action='store_true', help='Cho phép ghi đè bản xuất đã có')
    args = parser.parse_args()
    out = args.output.resolve()
    existing = [out / name for name in FILES if (out / name).exists()]
    if existing and not args.force:
        parser.error('File đã tồn tại; sao lưu rồi dùng --force hoặc chọn thư mục khác: ' + ', '.join(map(str, existing)))
    raw = json.loads(subprocess.check_output(['node', '-e', NODE_DATA], cwd=ROOT, text=True))
    enemies = source_rows('media/assets/enemy.xlsx', 'Enemy')
    npcs = source_rows('media/assets/npc.xlsx', 'NPC chính')
    with (ROOT / 'media/assets/chars/npc/Doi_Chieu_Ten_Anh.csv').open(encoding='utf-8-sig', newline='') as f:
        portraits = {r['npc_id']: r['new_path'] for r in csv.DictReader(f) if r['status'] == 'anh_npc_chinh'}
    data = build_data(raw, enemies, npcs, portraits)
    books = [make_book(data, [('hạt giống','Hạt giống'),('linh thảo','Linh thảo'),('đan dược','Đan dược'),
                               ('đan phương','Đan phương'),('pháp bảo','Pháp bảo'),('hiệu ứng pháp bảo','Hiệu ứng pháp bảo'),
                               ('công pháp','Công pháp'),('bí tịch','Bí tịch'),('rèn pháp bảo','Rèn pháp bảo'),('phù lục','Phù lục'),('vẽ phù','Vẽ phù')]),
             make_book(data, [('quái vật','Quái vật'),('boss','Boss'),('giai đoạn boss','Giai đoạn boss'),('vật phẩm rơi','Vật phẩm rơi')]),
             make_book(data, [('npc','NPC'),('quy tắc npc','Quy tắc NPC')]),
             make_book(data, [('linh căn','Linh căn'),('loại linh căn','Loại linh căn'),('ngũ hành','Ngũ hành'),('quy tắc linh căn','Quy tắc linh căn')])]
    out.mkdir(parents=True, exist_ok=True)
    for filename, wb in zip(FILES, books):
        wb.save(out / filename)
        print(out / filename)
    (out / FILES[-1]).write_text(json.dumps(data, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(out / FILES[-1])


if __name__ == '__main__':
    main()
