#!/usr/bin/env python3
"""Xuất bản tham chiếu Excel từ dữ liệu workspace; không cập nhật gameplay."""

import argparse
import csv
import json
import math
import subprocess
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parent.parent
DEFAULT_OUTPUT = ROOT / 'docs' / 'du-lieu-game'
FILES = ('vat_pham_game.xlsx', 'quai_va_boss_game.xlsx', 'npc_game.xlsx')

# Đây là đề xuất nội dung, không phải dữ liệu hay luật đang chạy trong game.
# Mỗi bộ gồm: bối cảnh, quan hệ với người chơi, thích, không thích.
NPC_IDEAS = {
    'NPC-01': ('Người sống tại Trấn Lạc Phong, quen các lối đi quanh trấn.', 'Người quen đầu hành trình; có thể chỉ đường.', 'Giúp người trong trấn; tin tức an toàn', 'Gây rối trong trấn'),
    'NPC-02': ('Tu sĩ Thanh Vân Tông chú trọng kiếm đạo và hộ sơn.', 'Đồng môn hoặc người hướng dẫn nhập môn.', 'Luận kiếm; bảo vệ đồng môn', 'Lạm dụng kiếm pháp'),
    'NPC-03': ('Tu sĩ Huyết Sát Môn theo đuổi sức mạnh qua chiến đấu.', 'Đối thủ hoặc đồng minh tùy lựa chọn.', 'Tỉ thí sòng phẳng; can đảm', 'Thất hứa; né tránh giao ước'),
    'NPC-04': ('Người của Vạn Dược Cốc nghiên cứu linh thảo và đan dược.', 'Người hướng dẫn hái thuốc và luyện đan.', 'Linh thảo; cứu người', 'Lãng phí dược liệu'),
    'NPC-05': ('Người tìm hiểu cơ quan ở di tích Huyền Quy.', 'Có thể trao manh mối thám hiểm.', 'Giải cơ quan; cổ thư', 'Phá hủy di tích'),
    'NPC-06': ('Người của Thiên Cơ Các quan sát biến động Linh Giới.', 'Nguồn tin về bí cảnh và dị tượng.', 'Thông tin chính xác; giữ bí mật', 'Tin đồn vô căn cứ'),
    'NPC-07': ('Tiên nhân gắn với Thiên Sổ Thành và việc ghi chép thiên mệnh.', 'Người giải thích thiên mệnh hoặc giao ủy thác.', 'Giữ lời hứa; bảo vệ thiên sổ', 'Can thiệp tùy tiện vào số mệnh'),
    'NPC-08': ('Nhân vật bí ẩn tại U Minh Cổ Điện, liên hệ chuyện luân hồi.', 'Người đối thoại trên đường chuyển thế.', 'Tôn trọng ký ức; lựa chọn có trách nhiệm', 'Xúc phạm người đã khuất'),
    'NPC-09': ('Người trông coi Đài Quy Hương tại Tiên Giới.', 'Người chỉ dẫn về đường trở về.', 'Giữ lời hẹn; tưởng nhớ cố nhân', 'Làm hại khách qua đường'),
    'NPC-10': ('Người quen đường trong Hắc Lâm, thường gặp yêu thú.', 'Bạn đồng hành trong ủy thác rừng.', 'Bảo vệ rừng; giúp người gặp nạn', 'Săn bắt vô tội vạ'),
    'NPC-11': ('Người làm việc quanh Mỏ Linh Thiết, am hiểu khoáng vật.', 'Người giới thiệu nguyên liệu rèn.', 'Linh thiết; dụng cụ tốt', 'Phá hoại mỏ'),
    'NPC-12': ('Thương nhân ở Lạc Phong Phường Thị.', 'Đối tác trao đổi hoặc nguồn tin thị trường.', 'Giao dịch công bằng; hàng hiếm', 'Gian lận'),
    'NPC-13': ('Tán tu tìm hiểu dòng linh khí ở Suối Địa Mạch.', 'Bạn luận đạo và khảo sát địa mạch.', 'Thanh tĩnh; nghiên cứu linh khí', 'Làm ô nhiễm linh tuyền'),
    'NPC-14': ('Lão nhân ẩn cư tại Kiếm Cốt Lâm, am hiểu kiếm ý.', 'Tiền bối có thể thử thách người chơi.', 'Kiên trì luyện kiếm; khiêm nhường', 'Khoe khoang võ lực'),
    'NPC-15': ('Tu sĩ thường đi qua Giới Nhai, quen đường giữa các giới.', 'Người dẫn đường khi thám hiểm.', 'Giúp khách lữ hành; phong cảnh', 'Bỏ rơi đồng hành'),
    'NPC-16': ('Người thăm dò khoáng vật ở Tinh Không Khoáng Vực.', 'Đối tác tìm linh tài.', 'Khoáng vật hiếm; khám phá', 'Phá hoại khu mỏ'),
    'NPC-17': ('Người nghiên cứu thảo mộc tại Bách Thảo Thánh Viện.', 'Người trao đổi kiến thức dược thảo.', 'Dược thảo; chăm sóc cây', 'Nhổ cây non tùy tiện'),
    'NPC-18': ('Lữ khách giấu danh tính quanh di tích Thiên Cơ.', 'Nhân vật bí ẩn có thể chia sẻ manh mối.', 'Giải bí ẩn; giữ kín danh tính', 'Truy hỏi đời tư'),
    'NPC-19': ('Tu sĩ tại Thành Vân Trung yêu thích vấn đạo.', 'Người luận đạo hoặc trao lời khuyên.', 'Đối thoại chân thành; trà', 'Áp đặt đạo lý'),
    'NPC-20': ('Người làm việc tại Thiên Sổ Thành, quan tâm cổ vật.', 'Người định giá hoặc trao đổi vật quý.', 'Cổ vật có nguồn gốc; giao dịch minh bạch', 'Hàng giả'),
    'NPC-21': ('Người chăm sóc cây cỏ tại Ngân Vân Tiên Viên.', 'Người hướng dẫn trồng tiên thảo.', 'Chăm vườn; hạt giống', 'Làm hỏng tiên viên'),
    'NPC-22': ('Lữ khách thường xuất hiện tại Đài Quy Hương.', 'Người đồng hành trong câu chuyện cố hương.', 'Kỷ vật; chuyện xưa', 'Phá hoại kỷ vật'),
    'NPC-23': ('Kiếm tu của Thái Hư Kiếm Cung.', 'Đối thủ luận kiếm hoặc tiền bối chỉ điểm.', 'Kiếm đạo; đấu công bằng', 'Đánh lén'),
    'NPC-24': ('Nhân vật tại Vô Tướng Ma Điện, nghiên cứu ma đạo.', 'Đối tác nguy hiểm hoặc đối thủ.', 'Kiến thức hiếm; giữ giao ước', 'Phản bội'),
}

# Các đặc điểm sau chỉ là gợi ý, không làm thay đổi luật tặng lễ hay PvP.
SOCIAL_IDEAS = {
    'trieu-nguyet': ('Kiếm tu lấy sự trong sáng làm phương châm.', 'Luận kiếm; giữ lời hẹn', 'Đánh lén'),
    'ho-thanh': ('Dược sư tán tu gắn bó với vườn thuốc.', 'Chăm linh thảo; cứu người', 'Phá vườn thuốc'),
    'bach-tuyet': ('Truyền nhân Băng cốc trân trọng những lời hẹn.', 'Cảnh tuyết; sự chân thành', 'Thất hứa'),
    'mac-kiem': ('Kiếm si đi khắp nơi tìm bạn luận kiếm.', 'Kiếm pháp mới; tỉ thí', 'Tự mãn'),
    'ly-van-du': ('Kiếm tu thích trao đổi đạo lý bên chén trà.', 'Trà; đàm luận kiếm đạo', 'Gây sự vô cớ'),
    'to-mong': ('Người tu đan đạo quan tâm linh khí.', 'Luyện đan; chia sẻ kinh nghiệm', 'Lãng phí dược liệu'),
    'han-tu': ('Trận pháp sư xem bảo hộ quan trọng như công kích.', 'Nghiên cứu trận pháp; bình tĩnh', 'Phá trận tùy tiện'),
    'diep-lan': ('Người nghiên cứu sự chuyển hóa âm dương.', 'Luận đạo; cân bằng', 'Cố chấp'),
    'van-thuong': ('Lôi tu chú trọng giữ tâm vững trước sấm sét.', 'Tôi luyện tâm cảnh; lôi pháp', 'Nóng vội'),
    'dao-quang': ('Tu sĩ tinh tú thường suy ngẫm về nguồn cội.', 'Ngắm sao; chuyện cố hương', 'Quên lời hứa'),
    'hanh-chu': ('Võ giả luyện quyền, tìm cơ hội tỉ thí.', 'Đấu công bằng; luyện quyền', 'Đánh lén'),
    'lieu-yen': ('Kiếm tu theo đuổi kiếm pháp nhẹ như mây.', 'Luyện kiếm; đối thủ xứng tầm', 'Gian lận'),
    'thach-cuong': ('Thể tu rèn luyện thân thể bền bỉ.', 'Tu luyện; tinh thần kiên trì', 'Coi thường đối thủ'),
    'tich-da': ('Tu sĩ ám ảnh đạo thích những trận đấu tính toán.', 'Chiến thuật; đối thủ bình tĩnh', 'Phá bỏ quy tắc tỉ thí'),
}

NODE_DATA = r"""
const items = require('./shared/data/items.js');
const catalog = require('./shared/data/catalog.js');
const adventure = require('./shared/data/adventures.js');
const social = require('./shared/data/social.js');
const F = require('./shared/rules/formulas.js');
const Game = require('./shared/game.js').Game;
require('./client/assets/assets.js');
const A = globalThis.HU_ASSETS;
const byId = Object.fromEntries(items.map(x => [x.code, x]));
const enriched = items.map(x => {
  const y = {...x};
  if (x.kind === 'SEED') {
    if (!byId[x.herb]) throw Error('Missing herb for ' + x.code);
    y.growthSeconds = Math.max(60, F.craft_time(byId[x.herb].tier) * 2);
  }
  if (x.kind === 'EQ') {
    const flat = F.item_flat_bonus(x.tier, x.sub);
    y.bonus = {hp: F.P.HP0 * flat, atk: F.P.ATK0 * flat, def: F.P.DEF0 * flat};
  }
  if (x.kind === 'BT') y.bonus = F.cong_phap_effect(x.effect, x.sub, 0, 1);
  return y;
});
const combat = catalog.ZONES.flatMap(z => Object.keys(F.MONSTER_MULT).map(kind => ({
  zone: z, name: catalog.ENEMY_NAMES[z.id], kind, mult: F.MONSTER_MULT[kind],
  dir: A.enemyDir(z.id, kind), path: A.enemyPath(z.id, kind)
})));
const game = new Game();
const recipes = items.filter(x => x.kind === 'PILL').map(x => game.recipe(x.code));
console.log(JSON.stringify({items: enriched, combat, dungeons: adventure.DUNGEONS,
  social: {FRIENDS: social.FRIENDS, DAO: social.DAO, RIVALS: social.RIVALS},
  realms: F.REALMS, recipes}));
"""


def read_source_sheet(path, sheet):
    wb = load_workbook(ROOT / path, read_only=True, data_only=True)
    try:
        return list(wb[sheet].values)[6:]
    finally:
        wb.close()


def table(wb, name, headers, rows, percent_headers=()):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for row in rows:
        if len(row) != len(headers):
            raise ValueError(f'{name}: hàng có {len(row)} cột, cần {len(headers)}')
        ws.append(row)
    ws.freeze_panes = 'C2'
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    ws.row_dimensions[1].height = 29
    for cell in ws[1]:
        cell.fill = PatternFill('solid', fgColor='27485B')
        cell.font = Font(name='Calibri', bold=True, color='FFFFFF')
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name='Calibri', size=11)
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if cell.row % 2 == 0:
                cell.fill = PatternFill('solid', fgColor='F2F7F8')
            if headers[cell.column - 1] in percent_headers and isinstance(cell.value, (int, float)):
                cell.number_format = '0.00%'
            elif isinstance(cell.value, (int, float)) and abs(cell.value) >= 1e9:
                cell.number_format = '0.00E+00'
    for col, header in enumerate(headers, 1):
        values = [str(header)] + [str(row[col - 1] or '') for row in rows]
        width = min(50, max(12, max(map(len, values), default=12) + 2))
        ws.column_dimensions[get_column_letter(col)].width = width
    for row_number, row in enumerate(rows, 2):
        lines = max((math.ceil(len(str(value)) / (ws.column_dimensions[get_column_letter(col)].width - 2))
                     for col, value in enumerate(row, 1) if value is not None), default=1)
        ws.row_dimensions[row_number].height = min(90, max(18, lines * 15))
    return ws


def workbook(notes):
    wb = Workbook()
    wb.remove(wb.active)
    table(wb, 'Hướng dẫn', ['Chủ đề', 'Giải thích'], notes)
    wb['Hướng dẫn'].column_dimensions['A'].width = 26
    wb['Hướng dẫn'].column_dimensions['B'].width = 100
    return wb


def make_items(data, item_art):
    items = data['items']
    by_id = {x['code']: x for x in items}
    seeds = {x['herb']: x['code'] for x in items if x['kind'] == 'SEED'}
    wb = workbook([
        ('Phạm vi', 'Danh mục tham chiếu từ shared/data/items.js và các luật hiện tại; không tự thay đổi dữ liệu game.'),
        ('Tên và ảnh', 'Tên game dùng items.js; tên/ô ảnh theo media/assets/vat_pham.xlsx. Hai tên có thể khác nhau, cùng tra cứu bằng mã.'),
        ('Độ khó', 'Đề xuất theo bậc cảnh giới: 0–4 dễ, 5–8 vừa, 9–13 khó. Game chưa có chỉ số độ khó trồng riêng.'),
        ('Thời gian trồng', 'Giây thực, trước khi dùng phân bón: max(60, craft_time(bậc LINH THẢO) × 2). Cây chín vẫn chờ thu hoạch.'),
        ('Phẩm chất', 'Các số tính theo phẩm chất mặc định sub của item. Bản @0–@3 có thể khác; sub 0/1/2/3 = Hạ/Trung/Thượng/Cực.'),
        ('Pháp bảo', 'HP/ATK/DEF là cộng thẳng cơ sở trước các hệ số nhân khác, không phải tổng chỉ số nhân vật. Ô trống là không có cộng riêng trong luật.'),
        ('Hiệu ứng pháp bảo', 'Giáp ảo, hồi máu, đóng băng, độc, chảy máu, choáng chưa được gán cho EQ trong luật hiện tại; không tự suy đoán.'),
        ('Công pháp', 'Chỉ số % tính ở cùng bậc cảnh giới và lĩnh ngộ 100%; người chơi mới học được 5%, chênh lệch bậc và các trần khác làm thay đổi hiệu quả.'),
        ('Bí tịch', 'Game dùng chung loại BT cho công pháp và bí tịch. Chưa có BT riêng cho rèn pháp bảo/vẽ phù; hai hàng trong sheet là chỗ trống đề xuất.'),
        ('Ô trống', 'Ô trống nghĩa là không có dữ liệu hoặc không có hiệu ứng riêng, không phải một giá trị đã thiết kế.'),
    ])
    def art(x):
        a = item_art.get(x['code'], ())
        return [a[1] if a else None, a[3] if a else None, a[4] if a else None]

    rows = []
    for x in sorted((x for x in items if x['kind'] == 'SEED'), key=lambda x: x['code']):
        herb = by_id[x['herb']]
        difficulty = 'Dễ' if herb['tier'] <= 4 else 'Vừa' if herb['tier'] <= 8 else 'Khó'
        rows.append([x['code'], x['name'], x['herb'], herb['name'], x['map'], x['realm'],
                     'Mua tại cửa hàng; kỳ ngộ chọn theo bậc và giá',
                     difficulty, round(x['growthSeconds'], 3), x['description'], *art(x)])
    table(wb, 'Hạt giống', ['Mã hạt', 'Tên hạt trong game', 'Mã cây', 'Tên cây trong game', 'Giới', 'Cảnh giới',
                            'Nguồn nhận', 'Độ khó (đề xuất)', 'Thời gian trồng (giây)', 'Công dụng', 'Tên theo ảnh', 'Ảnh atlas', 'Ô atlas'], rows)

    rows = []
    for x in sorted((x for x in items if x['kind'] == 'HERB'), key=lambda x: x['code']):
        used_by = [r['code'] for r in data['recipes'] if x['code'] in r['herbs']]
        rows.append([x['code'], x['name'], x['map'], x['realm'],
                     'Mua; nhặt từ quái theo bậc; kỳ ngộ' + ('; trồng từ hạt' if x['code'] in seeds else ''),
                     'Bán tại cửa hàng' + ('; nguyên liệu luyện đan' if used_by else ''),
                     ', '.join(used_by), x['description'], x.get('age'), seeds.get(x['code']), *art(x)])
    table(wb, 'Linh thảo', ['Mã cây', 'Tên cây trong game', 'Giới', 'Cảnh giới', 'Nguồn nhận', 'Cách dùng trong game',
                            'Mã đan dùng cây', 'Mô tả', 'Tuổi dược liệu', 'Mã hạt', 'Tên theo ảnh', 'Ảnh atlas', 'Ô atlas'], rows)

    rows = []
    for x in sorted((x for x in items if x['kind'] == 'EQ'), key=lambda x: x['code']):
        b = x['bonus']
        rows.append([x['code'], x['name'], x['grade'], x['quality'], x['realm'], 'Trang bị: cộng HP, ATK và DEF',
                     b['hp'], None, b['atk'], b['def'], None, None, None, None, *art(x)])
    table(wb, 'Pháp bảo', ['Mã PB', 'Tên pháp bảo trong game', 'Độ hiếm (phẩm cấp)', 'Phẩm chất mặc định', 'Cảnh giới', 'Công dụng hiện tại',
                            'HP cộng', 'MP cộng', 'ATK cộng', 'DEF cộng', 'EVA cộng', 'STA cộng', 'CRIT cộng', 'CRIT_DMG cộng',
                            'Tên theo ảnh', 'Ảnh atlas', 'Ô atlas'], rows)

    rows = []
    manuals = sorted((x for x in items if x['kind'] == 'BT'), key=lambda x: x['code'])
    for x in manuals:
        b = x['bonus']
        rows.append([x['code'], x['name'], x['grade'], x['quality'], x['realm'], x['effect'], x['description'],
                     b.get('tang_HP_DEF'), None, b.get('tang_ATK'), b.get('tang_HP_DEF'), None, None, None, None,
                     b.get('cong_gop_toc_do'), b.get('ky_nang_luyen_dan'), *art(x)])
    pct = ('HP %', 'ATK %', 'DEF %', 'Tốc tu %')
    table(wb, 'Công pháp', ['Mã CP', 'Tên công pháp trong game', 'Độ hiếm (phẩm cấp)', 'Phẩm chất mặc định', 'Cảnh giới', 'Mã hiệu ứng', 'Mô tả',
                            'HP %', 'MP %', 'ATK %', 'DEF %', 'EVA %', 'STA %', 'CRIT %', 'CRIT_DMG %', 'Tốc tu %',
                            'Kỹ năng luyện đan', 'Tên theo ảnh', 'Ảnh atlas', 'Ô atlas'], rows, pct)
    types = {'TOC': 'Luyện khí', 'DAN': 'Luyện đan', 'SAT': 'Công pháp công kích', 'PHONG': 'Công pháp phòng ngự'}
    rows = [[x['code'], x['name'], types[x['effect']], x['grade'], x['quality'], x['description'], 'Có trong game'] for x in manuals]
    rows.extend([
        [None, None, 'Luyện pháp bảo', None, None, 'Đề xuất bổ sung bí tịch riêng; rèn hiện chưa yêu cầu bí tịch.', 'Đề xuất, chưa có'],
        [None, None, 'Vẽ phù lục', None, None, 'Đề xuất bổ sung bí tịch riêng; chế phù hiện chưa yêu cầu bí tịch.', 'Đề xuất, chưa có'],
    ])
    table(wb, 'Bí tịch', ['Mã bí tịch', 'Tên bí tịch trong game', 'Thể loại', 'Độ hiếm (phẩm cấp)', 'Phẩm chất mặc định', 'Công dụng', 'Trạng thái'], rows)

    rows = []
    for r in sorted(data['recipes'], key=lambda r: r['code']):
        x = by_id[r['code']]
        ingredients = '; '.join(f"{code} ×{count}" for code, count in r['herbs'].items())
        rows.append([x['code'], x['name'], x['realm'], ingredients, r['stone'], round(r['time'], 3), x['effect']])
    table(wb, 'Đan phương', ['Mã đan', 'Tên đan', 'Cảnh giới đan', 'Nguyên liệu', 'Phí linh thạch Hạ', 'Thời gian (giây)', 'Mã hiệu ứng'], rows)

    rows = []
    for x in sorted((x for x in items if x['kind'] not in ('SEED', 'HERB', 'EQ', 'BT')), key=lambda x: x['code']):
        rows.append([x['code'], x['name'], x['category'], x['kind'], x['map'], x['realm'], x['quality'],
                     x.get('effect'), x['description'], *art(x)])
    table(wb, 'Vật phẩm khác', ['Mã', 'Tên trong game', 'Nhóm', 'Loại game', 'Giới', 'Cảnh giới', 'Phẩm chất mặc định',
                                'Mã hiệu ứng', 'Mô tả game', 'Tên theo ảnh', 'Ảnh atlas', 'Ô atlas'], rows)
    return wb


def make_enemies(data, assets):
    wb = workbook([
        ('Phạm vi', 'Quái/boss theo media/assets/enemy.xlsx và đối chiếu luật lịch luyện; tên hiển thị trong game theo địa điểm, không theo tên ảnh.'),
        ('Mã và tên', 'Mã MON/BOSS là mã tài nguyên, không phải ID combat. Một tên game tại mỗi địa điểm được dùng cho cả 4 loại đối thủ.'),
        ('Ảnh', 'Đường dẫn trong file nguồn là đường dẫn cũ; đường dẫn hiện tại chỉ điền khi client có ánh xạ tới mã ảnh đó.'),
        ('Chỉ số', 'HP/ATK/DEF quái = stats(zone, min(tầng người chơi - 1, tầng tối đa zone - 1)) × hệ số loại; không có chỉ số cố định theo mã ảnh.'),
        ('Bí cảnh', 'Tên boss bí cảnh và mã bí cảnh lấy từ adventures.js; chưa có ID boss hoặc liên kết ảnh riêng, không tự gán mã.'),
        ('Thủ lĩnh khác giới', 'Hai boss Hỗn Độn trong file ảnh nguồn chưa được ánh xạ riêng trong client; không giả định đường dẫn hình hay chỉ số của chúng.'),
    ])
    mapped = {}
    for x in data['combat']:
        mapped['monster/' + x['dir']] = x['path']
    asset_by_path = {r[7]: r for r in assets if r and r[0]}
    for title, is_boss in [('Quái vật', False), ('Boss', True)]:
        rows = []
        for r in assets:
            if not r or not r[0] or (r[3] != 'Enemy') != is_boss:
                continue
            path = mapped.get(r[7])
            rows.append([r[0], r[1], r[2], r[3], 'Có ánh xạ lịch luyện' if path else 'Chưa có ánh xạ lịch luyện',
                         r[7], path, r[4], r[5], r[6], r[8]])
        table(wb, title, ['Mã ảnh', 'Tên theo ảnh', 'Giới ảnh', 'Loại ảnh', 'Sử dụng trong client', 'Đường dẫn nguồn cũ',
                          'Đường dẫn client hiện tại', 'Frame đứng', 'Frame tấn công', 'Frame tử vong', 'Mã ảnh cũ'], rows)
    rows = []
    for x in data['combat']:
        a = asset_by_path.get('monster/' + x['dir'])
        if not a:
            raise ValueError('Không tìm thấy mã ảnh cho ' + x['dir'])
        z = x['zone']
        rows.append([z['id'], z['name'], data['realms'][z['realm']], x['name'], x['kind'], x['mult'],
                     a[0], a[1], x['path']])
    table(wb, 'Lịch luyện', ['ID địa điểm', 'Địa điểm', 'Cảnh giới', 'Tên game', 'Loại đối thủ', 'Hệ số HP ATK DEF',
                             'Mã ảnh', 'Tên theo ảnh', 'Đường dẫn ảnh hiện tại'], rows)
    rows = [[d['id'], d['name'], d['boss'], data['realms'][d['realm']], d['floors'], d['seconds'], d['desc'],
             'Mô phỏng sức mạnh theo tầng; chưa có mã boss/ảnh riêng'] for d in data['dungeons']]
    table(wb, 'Boss bí cảnh', ['Mã bí cảnh', 'Bí cảnh', 'Tên boss', 'Cảnh giới mở', 'Số tầng', 'Thời gian (giây)', 'Bối cảnh', 'Ghi chú'], rows)
    return wb


def make_npcs(data, npcs, portraits):
    wb = workbook([
        ('Phạm vi', 'NPC chính lấy mã/tên/giới/địa điểm từ media/assets/npc.xlsx; đường dẫn ảnh hiện tại theo bảng đối chiếu trong chars/npc.'),
        ('Nội dung đề xuất', 'Bối cảnh, quan hệ, thích và không thích ở NPC chính là ý tưởng mới để duyệt, chưa có luật gameplay.'),
        ('Tình trạng', '24 NPC chính có trong danh mục ảnh; chưa tìm thấy định nghĩa gameplay tương ứng trong shared/data.'),
        ('Nhân vật mô phỏng', 'Tri kỷ/đạo hữu/đối thủ lấy ID và dữ liệu từ shared/data/social.js, là các nhân vật khác với 24 NPC ảnh; bối cảnh/thích/không thích là đề xuất.'),
        ('Tặng quà', 'Tặng lễ tri kỷ hiện dùng linh thạch và cộng tình nghĩa theo cảnh giới; không kiểm tra danh sách thích/không thích.'),
        ('Ảnh', 'Đường dẫn tính từ gốc workspace; cột đường dẫn trong npc.xlsx cũ là ui/npc, ảnh hiện tại nằm tại chars/npc.'),
    ])
    rows = []
    for r in npcs:
        if not r or not r[0]:
            continue
        code, name, world, location, _ = r[:5]
        if code not in NPC_IDEAS or code not in portraits:
            raise ValueError('Thiếu đề xuất hoặc đối chiếu ảnh NPC: ' + code)
        image = 'media/assets/chars/npc/' + portraits[code]
        if not (ROOT / image).is_file():
            raise ValueError('Không tìm thấy ảnh NPC: ' + image)
        rows.append([code, name, world, location, *NPC_IDEAS[code], 'Đề xuất; chưa có luật NPC', image])
    table(wb, 'NPC chính', ['Mã NPC', 'Tên NPC', 'Giới', 'Địa điểm', 'Bối cảnh (đề xuất)', 'Quan hệ (đề xuất)',
                            'Thích (đề xuất)', 'Không thích (đề xuất)', 'Trạng thái', 'Ảnh hiện tại'], rows)
    rows = []
    for group, label in [('FRIENDS', 'Tri kỷ'), ('DAO', 'Đạo hữu'), ('RIVALS', 'Đối thủ PvP')]:
        for x in data['social'][group]:
            if x['id'] not in SOCIAL_IDEAS:
                raise ValueError('Thiếu đề xuất cho nhân vật mô phỏng: ' + x['id'])
            rows.append([x['id'], x['name'], label, x['gender'], data['realms'][x['realm']], x.get('title') or x.get('school'),
                         x.get('line'), *SOCIAL_IDEAS[x['id']], 'Nhân vật mô phỏng; nội dung đề xuất'])
    table(wb, 'Nhân vật mô phỏng', ['ID game', 'Tên NPC', 'Quan hệ trong game', 'Giới tính', 'Cảnh giới', 'Danh hiệu hoặc đạo',
                                   'Lời thoại', 'Bối cảnh (đề xuất)', 'Thích (đề xuất)', 'Không thích (đề xuất)', 'Trạng thái'], rows)
    return wb


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', type=Path, default=DEFAULT_OUTPUT, help='Thư mục xuất (mặc định docs/du-lieu-game)')
    parser.add_argument('--force', action='store_true', help='Cho phép ghi đè các file Excel đã có')
    args = parser.parse_args()
    output = args.output.resolve()
    existing = [output / name for name in FILES if (output / name).exists()]
    if existing and not args.force:
        parser.error('Các file đã tồn tại; sao lưu trước và dùng --force, hoặc chọn --output khác: ' + ', '.join(map(str, existing)))
    data = json.loads(subprocess.check_output(['node', '-e', NODE_DATA], cwd=ROOT, text=True))
    item_rows = read_source_sheet('media/assets/vat_pham.xlsx', 'Vật phẩm')
    item_art = {r[0]: r for r in item_rows if r and r[0]}
    enemies = read_source_sheet('media/assets/enemy.xlsx', 'Enemy')
    npcs = read_source_sheet('media/assets/npc.xlsx', 'NPC chính')
    with (ROOT / 'media/assets/chars/npc/Doi_Chieu_Ten_Anh.csv').open(encoding='utf-8-sig', newline='') as f:
        portraits = {r['npc_id']: r['new_path'] for r in csv.DictReader(f) if r['status'] == 'anh_npc_chinh'}
    books = [make_items(data, item_art), make_enemies(data, enemies), make_npcs(data, npcs, portraits)]
    output.mkdir(parents=True, exist_ok=True)
    for name, wb in zip(FILES, books):
        path = output / name
        wb.save(path)
        print(path)


if __name__ == '__main__':
    main()
